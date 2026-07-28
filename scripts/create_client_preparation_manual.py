from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output/client-preparation-manual"
IMG = OUT / "screenshots"
DOCX = OUT / "客户交付前置准备操作手册.docx"

GREEN = "2F855A"
DEEP = "1F2937"
MUTED = "5B6472"
LIGHT = "EEF7F1"
LINE = "D9E4DD"
ORANGE = "B45309"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=DEEP, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Noto Sans CJK SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    run._element.rPr.rFonts.set(qn("w:cs"), "Noto Sans CJK SC")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text, url, color="2563EB"):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    r_pr.append(color_node)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    run.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run2 = paragraph.add_run(" 页")
    run2.font.size = Pt(9)


def style_run(run, size=10.5, bold=False, color=DEEP):
    run.font.name = "Noto Sans CJK SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    run._element.rPr.rFonts.set(qn("w:cs"), "Noto Sans CJK SC")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        style_run(p.add_run(bold_prefix), bold=True)
        style_run(p.add_run(text[len(bold_prefix):]))
    else:
        style_run(p.add_run(text))
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.22)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    style_run(p.add_run(text), size=10.3)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    style_run(p.add_run(text), size=10.3)
    return p


def add_callout(doc, title, text, fill=LIGHT, accent=GREEN):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    style_run(p.add_run(title), bold=True, color=accent, size=11)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    style_run(p2.add_run(text), size=10.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_screenshot(doc, filename, caption, width=6.25):
    path = IMG / filename
    if not path.exists():
        add_callout(doc, "截图缺失", f"未找到 {filename}，请重新采集后补入。", fill="FFF7ED", accent=ORANGE)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(7)
    style_run(cap.add_run(caption), size=9, color=MUTED)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        set_cell_text(header.cells[i], text, bold=True, color="FFFFFF", size=10)
        set_cell_shading(header.cells[i], GREEN)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], str(text), size=9.7)
            if len(table.rows) % 2 == 0:
                set_cell_shading(cells[i], "F7FAF8")
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_link_line(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    style_run(p.add_run(label + "："), bold=True, size=10.3)
    add_hyperlink(p, url, url)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Noto Sans CJK SC"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
styles["Normal"]._element.rPr.rFonts.set(qn("w:cs"), "Noto Sans CJK SC")
styles["Normal"].font.size = Pt(10.5)
for name, size, color in (("Heading 1", 17, GREEN), ("Heading 2", 13.5, DEEP), ("Heading 3", 11.5, GREEN)):
    st = styles[name]
    st.font.name = "Noto Sans CJK SC"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    st._element.rPr.rFonts.set(qn("w:cs"), "Noto Sans CJK SC")
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = True

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
style_run(fp.add_run("客户交付前置准备操作手册 | 秒哒企业版 + 微信公众平台"), size=8.5, color=MUTED)
add_page_number(footer.add_paragraph())

# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(70)
style_run(p.add_run("客户交付前置准备操作手册"), size=25, bold=True, color=GREEN)
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
style_run(p2.add_run("秒哒企业版开通与微信公众平台（小程序）配置"), size=14, color=DEEP)
p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
p3.paragraph_format.space_before = Pt(18)
style_run(p3.add_run("适用于：AI 营养秤微信小程序交付前的客户准备工作"), size=10.5, color=MUTED)
p4 = doc.add_paragraph()
p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
p4.paragraph_format.space_before = Pt(100)
style_run(p4.add_run("版本：2026-07-27"), size=10, color=MUTED)
doc.add_page_break()

add_heading(doc, "一、先看结论", 1)
add_callout(doc, "本手册的边界", "当前交付物是微信小程序，不是公众号网页应用。客户主要操作微信公众平台中的“小程序后台”；独立公众号不是本项目运行的硬性前置条件。若客户已有公众号，可按本手册末尾的“公众号关联（可选）”操作。")
add_body(doc, "秒哒企业版目前通过官网的“企业版咨询”入口提交企业信息，由百度智能云销售/商务团队跟进报价、合同和企业空间开通；官网没有面向客户的自助付款按钮。")
add_body(doc, "客户完成本手册后，应向交付方提供：秒哒企业空间管理员信息、小程序 AppID、AppSecret（安全传输）、合法域名、隐私政策主体信息，以及订阅消息模板 ID 和字段截图。")

add_heading(doc, "二、客户前期准备清单", 1)
add_table(doc, ["事项", "客户准备内容", "交付方使用方式"], [
    ("秒哒企业版", "企业名称、联系人、手机号、邮箱、企业规模、行业、预计席位/场景", "开通企业空间、团队成员和项目权限"),
    ("小程序主体", "营业执照、管理员身份证/微信、企业邮箱、名称、头像、简介、服务类目", "上传代码、配置 AppID、提交审核"),
    ("小程序凭证", "AppID、AppSecret", "后端微信登录、Access Token、订阅消息发送"),
    ("域名", "已备案 HTTPS 域名及 DNS/证书管理人", "小程序 request/download/socket 合法域名"),
    ("隐私材料", "主体名称、联系方式、隐私政策、第三方服务说明、数据保存期限", "小程序隐私保护指引和真机验收"),
    ("订阅消息", "早餐、午餐、晚餐、饮水四类模板 ID 及字段编号", "配置提醒模板并联调发送"),
], [1.25, 3.0, 2.05])

add_heading(doc, "三、步骤 1：购买并开通秒哒企业版账户", 1)
add_heading(doc, "3.1 进入秒哒官网并登录/注册", 2)
add_body(doc, "打开秒哒官网，使用企业管理员的百度账号登录；没有账号时先按页面提示注册。此处只完成账号准备，不要把企业管理员账号交给第三方代操作。")
add_link_line(doc, "秒哒官网", "https://www.miaoda.cn/")
add_screenshot(doc, "01_miaoda_home.png", "图 1-1 秒哒官网首页：右上角可见“企业版咨询”入口。")

add_heading(doc, "3.2 提交企业版咨询", 2)
add_body(doc, "在秒哒首页点击“企业版咨询”，进入百度智能云的企业版信息咨询表单。填写公司名称、联系人姓名、联系人手机、联系人邮箱、企业规模、所属行业和使用场景等信息。")
add_link_line(doc, "企业版咨询/开通入口", "https://cloud.baidu.com/survey/miaoda_enterprise.html")
add_screenshot(doc, "02_miaoda_enterprise_consult.png", "图 1-2 企业版咨询表单：带 * 的字段为必填项；手机号和验证码由客户本人完成。")
add_callout(doc, "客户操作边界", "表单提交、验证码接收、商务沟通、报价确认、合同签署和付款均由客户完成。交付方可以协助确认技术场景和席位数量，但不代填客户的企业身份、手机号或付款信息。", fill="FFF7ED", accent=ORANGE)

add_heading(doc, "3.3 与百度智能云确认方案并完成开通", 2)
for text in [
    "向销售说明：需要开通秒哒企业版，预计企业成员数量、项目数量、协作方式，以及是否需要后端服务、数据隔离或专属支持。",
    "确认报价、服务范围、席位数量、续费方式、SLA 和数据/权限边界。",
    "完成合同和付款后，由百度智能云开通企业组织或企业空间，并指定企业管理员。",
    "企业管理员进入秒哒后，创建团队空间、邀请成员，并为交付人员分配开发和发布所需权限。",
]:
    add_number(doc, text)

add_heading(doc, "3.4 阅读秒哒企业版产品文档", 2)
add_body(doc, "企业版采用“企业 - 团队 - 成员”三级权限结构。客户管理员应先理解成员管理、团队空间、秒点/资源分配和项目权限，再邀请交付方进入对应项目。")
add_link_line(doc, "秒哒企业版官方产品文档", "https://cloud.baidu.com/doc/MIAODA/s/omnepp22a")
add_screenshot(doc, "03_miaoda_enterprise_docs.png", "图 1-3 百度智能云官方文档：秒哒企业版介绍与企业版结构。")

add_heading(doc, "四、步骤 2：微信公众平台配置", 1)
add_callout(doc, "先确认账号类型", "本项目需要的是“小程序”。客户已有公众号时，不需要为了运行小程序再单独注册一个公众号；公众号关联属于可选的品牌导流配置。")

add_heading(doc, "4.1 进入微信公众平台", 2)
add_body(doc, "打开微信公众平台，客户管理员可以选择扫码登录或账号登录。后续所有 AppID、AppSecret、开发设置、隐私保护指引、订阅消息和版本发布操作，都在该平台完成。")
add_link_line(doc, "微信公众平台", "https://mp.weixin.qq.com/")
add_screenshot(doc, "04_wechat_platform_login.png", "图 2-1 微信公众平台入口：使用账号登录或扫码登录。")
add_callout(doc, "后台截图说明", "AppID、AppSecret、合法域名、隐私保护指引、订阅消息和发布页面均需客户登录自己的微信公众平台账号后才能看到。出于账号安全、验证码和企业敏感信息边界，本手册不代客户登录后台；对应页面按后台菜单路径提供操作步骤，客户可登录后逐项完成。", fill="FFF7ED", accent=ORANGE)

add_heading(doc, "4.2 注册或确认小程序账号", 2)
add_body(doc, "首次准备时，点击“立即注册”，在账号类型页面选择“小程序”，不要选择公众号或服务号。")
add_screenshot(doc, "05_wechat_account_types.png", "图 2-2 注册账号类型选择：本项目选择“小程序”。")
add_screenshot(doc, "06_wechat_miniprogram_intro.png", "图 2-3 小程序注册说明：注册后还要完成开发、审核和发布。")
add_body(doc, "注册前请准备以下材料：")
for text in [
    "企业营业执照、法定代表人/管理员身份材料，以及用于微信验证的管理员微信。",
    "企业邮箱或可长期维护的账号邮箱；不要使用离职人员私人邮箱。",
    "小程序名称、头像、简介、服务类目和品牌资质。健康管理、AI 和智能硬件相关类目以微信审核结果为准。",
    "用于服务器域名的 HTTPS 域名。正式域名应完成备案并能稳定解析。",
]:
    add_bullet(doc, text)

add_heading(doc, "4.3 完成认证并添加开发者", 2)
for text in [
    "完成企业主体认证和小程序基本信息填写。",
    "进入“管理 - 成员管理”，添加交付方开发者和体验成员。",
    "进入“开发 - 开发管理”，记录 AppID；在开发设置中查看或生成 AppSecret。",
    "AppSecret 通过密码管理器、加密文件或一对一安全渠道提供给交付方，不要在群聊、工单或普通文档中明文发送。",
]:
    add_number(doc, text)

add_heading(doc, "4.4 配置服务器域名", 2)
add_body(doc, "进入“开发 - 开发管理 - 开发设置 - 服务器域名”，配置以下域名。域名只填写协议 + 域名，不填写具体接口路径。")
add_table(doc, ["域名类型", "当前项目使用", "用途"], [
    ("request 合法域名", "https://zhgdvfwemwcmnehoarwp.supabase.co", "登录、数据、Edge Functions、图片上传"),
    ("downloadFile 合法域名", "https://zhgdvfwemwcmnehoarwp.supabase.co", "头像和食材图片等云端资源"),
    ("socket 合法域名", "wss://rtc-aiotgw.exp.bcelive.com", "AI 语音和实时对话"),
    ("uploadFile 合法域名", "当前版本暂无硬性要求", "当前图片上传通过后端请求完成"),
], [1.45, 2.75, 2.1])
add_callout(doc, "生产环境注意", "当前 Supabase 域名和百度 RTC 网关是现有项目配置。正式上线前请确认域名备案、证书和生产网关合规；如更换为客户自有域名，需同步更新小程序配置和后端环境变量。", fill="FFF7ED", accent=ORANGE)

add_heading(doc, "4.5 配置用户隐私保护指引", 2)
add_body(doc, "进入“设置 - 服务内容声明/用户隐私保护指引”，根据客户主体真实情况填写收集的信息、使用目的、保存期限、客服联系方式和第三方服务。项目会使用以下能力：")
add_table(doc, ["信息/接口", "客户需要说明的用途"], [
    ("微信身份标识（OpenID，可能包含 UnionID）", "微信登录、账号绑定和提醒发送"),
    ("头像、昵称、邮箱和账号信息", "创建和维护用户账号及个人资料"),
    ("健康档案、饮食记录", "营养分析、统计和个性化建议"),
    ("相册、摄像头", "上传食材图片并进行识别"),
    ("麦克风/录音", "语音输入和 AI 问答"),
    ("蓝牙设备信息", "搜索智能秤并读取称重数据"),
    ("设备型号和系统信息", "适配不同终端和蓝牙设备"),
    ("百度 AI、百度语音、Supabase", "完成图像识别、语音处理和云端存储"),
], [2.1, 4.2])
add_body(doc, "客户还需要把小程序内的隐私政策中的主体名称、联系方式、更新时间和数据保存期限改成真实信息。")

add_heading(doc, "4.6 配置一次性订阅消息", 2)
add_body(doc, "进入“功能 - 订阅消息”，添加并审核以下四类一次性模板。审核完成后，把模板 ID 和关键词编号截图发给交付方。")
add_table(doc, ["提醒类型", "代码需要的字段", "说明"], [
    ("早餐提醒", "餐单、日期、打卡时间", "模板 ID 与关键词编号由客户后台生成"),
    ("午餐提醒", "餐单、日期、打卡时间", "模板 ID 与关键词编号由客户后台生成"),
    ("晚餐提醒", "餐单、日期、打卡时间", "模板 ID 与关键词编号由客户后台生成"),
    ("饮水提醒", "温馨提示、饮水时间", "模板 ID 与关键词编号由客户后台生成"),
], [1.4, 2.5, 2.4])
add_callout(doc, "交付方配置项", "交付方会在后端配置 WECHAT_SUBSCRIBE_TEMPLATES_JSON。客户不需要接触后端密钥，只需提供模板 ID、模板字段编号和审核结果。", fill=LIGHT, accent=GREEN)

add_heading(doc, "4.7 公众号关联（可选）", 2)
add_body(doc, "如果客户希望从已有公众号菜单进入小程序，登录公众号后台后进入“小程序管理/关联小程序”，填写小程序 AppID 并完成管理员确认；再在公众号自定义菜单中添加“跳转小程序”，页面路径使用 `pages/home/index`。")
add_body(doc, "当前项目没有公众号网页授权、公众号模板消息、公众号客服消息或微信支付代码，因此公众号关联不是本项目运行的必需条件。")

add_heading(doc, "4.8 提交审核与发布", 2)
for text in [
    "交付方使用客户 AppID 构建并上传体验版本。",
    "客户管理员在体验版中验收微信登录、邮箱账户绑定、图片上传、麦克风、蓝牙和订阅提醒。",
    "客户在小程序后台填写服务类目、版本说明、测试账号和审核说明，提交微信审核。",
    "审核通过后由客户管理员点击发布；发布后再用真实手机和真实微信账号复测。",
]:
    add_number(doc, text)

add_heading(doc, "五、交付方需要客户回传的资料", 1)
add_table(doc, ["资料", "回传方式", "备注"], [
    ("秒哒企业空间管理员/成员信息", "一对一确认", "用于邀请交付人员并分配项目权限"),
    ("小程序 AppID", "普通文本即可", "用于构建和校验账号一致性"),
    ("小程序 AppSecret", "密码管理器或加密渠道", "禁止群聊明文发送"),
    ("合法域名", "普通文本 + 证书/备案状态", "生产环境需确认可被微信接受"),
    ("订阅模板 ID + 字段截图", "截图或加密文件", "早餐/午餐/晚餐/饮水四类"),
    ("隐私政策主体资料", "Word/文本确认", "主体名称、联系方式、第三方服务和保存期限"),
    ("审核联系人和测试账号", "一对一确认", "用于提审和审核问题处理"),
], [2.0, 2.15, 2.15])

add_heading(doc, "六、当前项目不需要客户配置的内容", 1)
for text in [
    "微信支付：当前代码没有 wx.requestPayment、商户号或支付回调。",
    "公众号网页 OAuth：当前登录使用小程序 wx.login，不需要公众号网页授权回调域名。",
    "公众号模板消息/客服消息：当前只使用小程序一次性订阅消息。",
    "微信手机号快速验证：当前版本没有调用 getuserphonenumber，旧部署说明中的手机号额度要求不适用于当前版本。",
    "百度 AI 和 RTC 密钥：由交付方配置到后端，不由客户写入小程序代码。",
]:
    add_bullet(doc, text)

add_heading(doc, "七、上线前验收清单", 1)
add_table(doc, ["检查项", "结果"], [
    ("客户可以登录秒哒并进入企业空间", "□"),
    ("交付方已加入企业空间/项目并具备开发权限", "□"),
    ("小程序主体已认证，AppID/AppSecret 已确认", "□"),
    ("request、downloadFile、socket 合法域名已配置", "□"),
    ("用户隐私保护指引和应用内隐私政策已更新", "□"),
    ("四类一次性订阅消息模板已审核并提供字段编号", "□"),
    ("微信登录、图片、麦克风、蓝牙真机验收通过", "□"),
    ("客户已提交审核并完成正式发布", "□"),
], [5.7, 0.5])

add_heading(doc, "八、官方入口汇总", 1)
add_link_line(doc, "秒哒官网", "https://www.miaoda.cn/")
add_link_line(doc, "秒哒企业版咨询/开通", "https://cloud.baidu.com/survey/miaoda_enterprise.html")
add_link_line(doc, "秒哒企业版官方产品文档", "https://cloud.baidu.com/doc/MIAODA/s/omnepp22a")
add_link_line(doc, "微信公众平台", "https://mp.weixin.qq.com/")
add_link_line(doc, "微信小程序官方开发文档", "https://developers.weixin.qq.com/miniprogram/dev/framework/")

doc.core_properties.title = "客户交付前置准备操作手册"
doc.core_properties.subject = "秒哒企业版与微信公众平台小程序配置流程"
doc.core_properties.author = "交付团队"
doc.core_properties.comments = "Screenshots captured from official public entry pages on 2026-07-27."
OUT.mkdir(parents=True, exist_ok=True)
doc.save(DOCX)
print(DOCX)
